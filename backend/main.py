import logging
import os
from contextlib import asynccontextmanager
from functools import lru_cache
from pathlib import Path
from typing import AsyncIterator

from docx import Document
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from groq import APIError, Groq
from pydantic import BaseModel, Field

logger = logging.getLogger("uvicorn.error")

PROJECT_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIR = PROJECT_DIR / "frontend"
RESUME_PATH = Path(__file__).resolve().with_name("Atharva_Kulkarni_Resume_Starter.docx")
GITHUB_URL = "https://github.com/Atharva97l"

load_dotenv(PROJECT_DIR / ".env")
MODEL = (os.getenv("GROQ_MODEL") or "openai/gpt-oss-120b").strip()


def read_resume(path: Path) -> str:
    if not path.is_file():
        raise FileNotFoundError(f"Resume document not found: {path}")

    document = Document(path)
    content = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                content.append(" | ".join(cells))

    resume_text = "\n".join(content)
    if not resume_text:
        raise ValueError(f"No text could be extracted from the resume: {path}")
    return resume_text


@lru_cache(maxsize=1)
def get_groq_client() -> Groq:
    api_key = (os.getenv("GROQ_API_KEY") or "").strip().strip('"').strip("'")
    if not api_key:
        raise HTTPException(status_code=503, detail="GROQ_API_KEY is not configured.")
    return Groq(api_key=api_key)


@asynccontextmanager
async def lifespan(app: FastAPI) -> AsyncIterator[None]:
    app.state.resume_text = read_resume(RESUME_PATH)
    yield


app = FastAPI(
    title="AI Resume Chat",
    description="Ask questions grounded in Atharva Kulkarni's resume.",
    lifespan=lifespan,
)

# Serve the frontend assets from the same origin as the API.
app.mount("/static", StaticFiles(directory=FRONTEND_DIR), name="static")


class ChatRequest(BaseModel):
    question: str = Field(min_length=1, max_length=4000)


@app.get("/", include_in_schema=False)
def home() -> FileResponse:
    return FileResponse(FRONTEND_DIR / "index.html")


@app.get("/api/info")
def api_info() -> dict[str, str]:
    return {
        "message": "AI Resume Chat API",
        "github_url": GITHUB_URL,
        "chat_endpoint": "/chat",
    }


@app.post("/chat")
def chat(request: ChatRequest, http_request: Request) -> dict[str, str]:
    system_prompt = f"""
You are a professional assistant answering questions on behalf of the candidate.
Use only the resume below. Do not infer or invent experience, skills, education,
contact details, or repository details. If the resume does not contain an answer,
say: "I don't have enough information to answer that."

The candidate's GitHub profile is {GITHUB_URL}. This is a profile link only;
do not claim to know its contents.

Resume:
{http_request.app.state.resume_text}
"""
    try:
        response = get_groq_client().chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": request.question},
            ],
        )
    except HTTPException:
        raise  # e.g. 503 missing key; don't let the catch-all turn it into a 500
    except APIError as exc:
        logger.error("Groq API error: %s: %s | cause=%r", type(exc).__name__, exc, exc.__cause__)
        raise HTTPException(
            status_code=502,
            detail="The AI provider could not complete the request.",
        ) from exc
    except Exception:
        logger.exception("Unexpected error in /chat")
        raise HTTPException(status_code=500, detail="Unexpected server error.")

    answer = response.choices[0].message.content
    if not answer:
        raise HTTPException(status_code=502, detail="The AI provider returned an empty response.")

    return {"answer": answer}
        raise HTTPException(
            status_code=500,
            detail=f"{type(exc).__name__}: {str(exc)[:200]}",
        ) from exc
